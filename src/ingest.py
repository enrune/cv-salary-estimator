"""
Ingest: PDF / DOCX → plain text.

Proč dva backendy: zadání povoluje oba formáty CV; pdfplumber a python-docx jsou
dvě nejstabilnější knihovny pro každý z nich. Detekce formátu podle přípony,
fallback chybí záměrně — pokud někdo pošle .txt, ať to selže rychle a hlasitě.
"""


from __future__ import annotations
# pathlib.Path místo os.path — modernější API, čitelnější (.suffix vs os.path.splitext)
from pathlib import Path

# pdfplumber — vybráno pro layout-aware extrakci. Alternativy:
# PyPDF2 (rychlejší, ale horší u sloupců), pymupdf (rychlejší, ale GPL → licenční riziko)
import pdfplumber

# python-docx — de facto standard pro .docx; čte přímo XML uvnitř archivu
from docx import Document


def extract_text(path: str | Path) -> str:
    """
    Extrahuje plain text z CV. Podporuje PDF a DOCX.
    Vrací jednolitý string s newlines mezi sekcemi (LLM rozumí newlines jako oddělovačům).
    """
    # Path() konverze přijme str i Path — uživatelsky přívětivé
    p = Path(path)

    # exists() check před otevřením — pdfplumber by jinak hodil méně srozumitelnou chybu
    if not p.exists():
        # FileNotFoundError je standardní výjimka, IDE/uživatel ji zná
        raise FileNotFoundError(f"Soubor neexistuje: {p}")

    # .suffix vrací včetně tečky (".pdf"), .lower() pro case-insensitive porovnání
    suffix = p.suffix.lower()

    # Větvení podle formátu — explicitní mapping je čitelnější než dict[suffix → callable]
    if suffix == ".pdf":
        return _extract_pdf(p)
    if suffix == ".docx":
        return _extract_docx(p)

    # Fail fast — neznámý formát = okamžitě informativní chyba, ne tichá degradace
    raise ValueError(
        f"Nepodporovaný formát: {suffix}. Podporované: .pdf, .docx"
    )


def _extract_pdf(p: Path) -> str:
    """Privátní helper pro PDF — podtržítko je Python konvence pro 'interní'."""
    # Context manager (with) — zajistí zavření file handle i při výjimce
    with pdfplumber.open(p) as pdf:
        # List comprehension přes všechny stránky.
        # `or ""` — extract_text() může vrátit None u prázdné/obrázkové stránky,
        # None by zlomil "\n".join, takže prázdný string je bezpečný fallback
        pages = [page.extract_text() or "" for page in pdf.pages]

    # Spojíme stránky newlinem; \n je univerzální oddělovač který LLM dobře parsuje
    return "\n".join(pages).strip()  # strip() odstraní wrapping whitespace pro čistší prompt


def _extract_docx(p: Path) -> str:
    """Privátní helper pro DOCX."""
    # Document() načte celý soubor do paměti — pro CV (max desítky KB) zanedbatelné
    doc = Document(str(p))  # str(p) protože Document neumí Path v některých verzích

    # paragraphs zahrnuje běžné odstavce; nadpisy a odrážky jsou taky paragraphs
    # Tabulky řešíme samostatně níže, protože jsou v doc.tables, ne paragraphs
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

    # Extrakce z tabulek — CV často mají tabulkový layout (zkušenosti, dovednosti)
    # Bez tohoto by se ztratil obsah celých tabulek
    table_texts: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            # Cell.text spojí všechny paragraphs v buňce; " | " jako oddělovač buněk
            # protože newline by ztratil řádkovou strukturu
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:  # Skip prázdných řádků — šum v promptu škodí přesnosti LLM
                table_texts.append(row_text)

    # Spojení paragraphs + tables; pořadí: paragraphs první, protože obvykle obsahují hlavní text
    return "\n".join(paragraphs + table_texts).strip()
